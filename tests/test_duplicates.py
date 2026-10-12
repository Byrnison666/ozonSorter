"""Дубли: выданная посылка (уже попала в выгрузку «Отгрузки») пропала из отчёта
склада и снова пришла с тем же номером и штрихкодом. В пределах DUPLICATE_DAYS —
дубль, в отгрузку до решения оператора не идёт; позже — новая посылка.
Пропуск только в неполном прошлом отчёте дублем не считается.
"""
import json
import os
import tempfile
import unittest
from datetime import timedelta

import openpyxl
from sqlalchemy import select

from src.database import DatabaseManager
from src.duplicate_service import (
    DUPLICATE_DAYS, DuplicateService, is_incomplete,
)
from src.export_service import ExportService
from src.models import AssignmentStatus, Client, DeliveryPoint, Shipment
from src.services import ImportService
from src.stale_service import StaleShipmentService

POINT = DeliveryPoint.KOMSOMOLSKAYA_4
OTHERS = [f"111-OTH-{i}" for i in range(1, 10)]


def _make_report(postings):
    path = tempfile.mktemp(suffix=".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Этикетка\nНазвание", "Номер отправления", "Тип", "Статус", "Ячейка"])
    for p in postings:
        ws.append([f"LBL\n{p}", p, "Обычный", "Готово к выдаче", "A-01"])
    wb.save(path)
    return path


class DuplicatesTest(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        self.db = DatabaseManager(db_path=self.db_path)
        self.db.create_tables()
        self.session = self.db.get_session()
        self.session.add(Client(
            ozon_client_id="111", full_name="Тест", fixed_delivery_point=POINT,
        ))
        self.session.commit()
        self.importer = ImportService(self.session)
        self._files = []

    def tearDown(self):
        self.session.close()
        os.remove(self.db_path)
        for f in self._files:
            if os.path.exists(f):
                os.remove(f)

    def _import(self, postings):
        path = _make_report(postings)
        self._files.append(path)
        return self.importer.process_import(path)

    def _export(self, import_session):
        """Выгрузка «Отгрузки» по отчёту — список номеров в файле."""
        out = tempfile.mktemp(suffix=".xlsx")
        self._files.append(out)
        ExportService(self.session).generate_export(POINT, out, import_session.id)
        ws = openpyxl.load_workbook(out).active
        return sorted(row[1] for row in ws.iter_rows(values_only=True) if len(row) > 1 and row[1])

    def _shipment(self, posting_number):
        return self.session.execute(
            select(Shipment).where(Shipment.posting_number == posting_number)
        ).scalar_one()

    def _age(self, posting_number, days):
        """Сдвинуть последнее появление посылки на days дней назад."""
        shipment = self._shipment(posting_number)
        shipment.last_seen_at -= timedelta(days=days)
        self.session.commit()

    def _issue_and_return(self, age_days=0):
        """A выгружена и выдана (пропала из отчёта), затем пришла снова."""
        day1 = self._import(["111-A-1"] + OTHERS)
        self._export(day1)
        self._import(OTHERS)
        if age_days:
            self._age("111-A-1", age_days)
        return self._import(["111-A-1", "111-NEW-1"] + OTHERS)

    def test_issued_shipment_reappearing_is_duplicate(self):
        day1 = self._import(["111-A-1"] + OTHERS)
        self._export(day1)
        issued_seen_at = self._shipment("111-A-1").last_seen_at
        self._import(OTHERS)
        day3 = self._import(["111-A-1", "111-NEW-1"] + OTHERS)

        a = self._shipment("111-A-1")
        self.assertEqual(a.assignment_status, AssignmentStatus.DUPLICATE)
        self.assertEqual(a.delivered_at, issued_seen_at)
        self.assertEqual(a.first_seen_at, a.last_seen_at)
        self.assertEqual(a.last_seen_import_session_id, day3.id)
        self.assertEqual(day3.new_to_ship_rows, 1)
        self.assertIn("DUPLICATE: Shipment 111-A-1 was issued and reappeared in import.",
                      json.loads(day3.log_json))
        self.assertEqual(self._export(day3), ["111-NEW-1"], "дубль не идёт в отгрузку")
        self.assertEqual(self.importer.skipped_duplicates, 0)

    def test_reappearing_on_the_last_day_is_duplicate(self):
        self._issue_and_return(age_days=DUPLICATE_DAYS)
        self.assertEqual(self._shipment("111-A-1").assignment_status,
                         AssignmentStatus.DUPLICATE)

    def test_reappearing_after_window_is_new_shipment(self):
        day3 = self._issue_and_return(age_days=DUPLICATE_DAYS + 1)

        a = self._shipment("111-A-1")
        self.assertEqual(a.assignment_status, AssignmentStatus.TO_SHIP)
        self.assertIsNone(a.exported_import_session_id)
        self.assertIsNone(a.delivered_at)
        self.assertEqual(a.first_seen_at, a.last_seen_at)
        self.assertEqual(day3.new_to_ship_rows, 2)
        self.assertIn(f"Shipment 111-A-1 reappeared {DUPLICATE_DAYS + 1} days after issue, "
                      f"shipping as new.", json.loads(day3.log_json))
        self.assertEqual(self._export(day3), ["111-A-1", "111-NEW-1"])

    def test_exported_shipment_still_in_report_is_not_duplicate(self):
        day1 = self._import(["111-A-1"])
        self._export(day1)
        day2 = self._import(["111-A-1"])
        a = self._shipment("111-A-1")
        self.assertEqual(a.assignment_status, AssignmentStatus.TO_SHIP)
        self.assertEqual(a.exported_import_session_id, day1.id)
        self.assertEqual(json.loads(day2.log_json), [])

    def test_not_exported_shipment_reappearing_is_not_duplicate(self):
        self._import(["111-A-1"] + OTHERS)
        self._import(OTHERS)
        day3 = self._import(["111-A-1"] + OTHERS)
        a = self._shipment("111-A-1")
        self.assertEqual(a.assignment_status, AssignmentStatus.TO_SHIP)
        self.assertEqual(self._export(day3), ["111-A-1"] + OTHERS)

    def test_list_and_ship_duplicate(self):
        day3 = self._issue_and_return()
        service = DuplicateService(self.session)
        rows = service.list_duplicates()
        self.assertEqual([r.posting_number for r in rows], ["111-A-1"])
        row = rows[0]
        self.assertEqual(row.client_name, "Тест")
        self.assertEqual(row.point, POINT)
        self.assertEqual(row.cell, "A-01")

        self.assertTrue(service.ship(row.shipment_id))
        self.assertFalse(service.ship(row.shipment_id), "уже не дубль")
        a = self._shipment("111-A-1")
        self.assertEqual(a.assignment_status, AssignmentStatus.TO_SHIP)
        self.assertIsNone(a.exported_import_session_id)
        self.assertEqual(service.list_duplicates(), [])
        self.assertEqual(self._export(day3), ["111-A-1", "111-NEW-1"])

    def test_dismissed_duplicate_is_silent_until_it_returns_again(self):
        self._issue_and_return()
        service = DuplicateService(self.session)
        self.assertTrue(service.dismiss(self._shipment("111-A-1").id))
        self.assertEqual(self._shipment("111-A-1").assignment_status,
                         AssignmentStatus.DELIVERED)
        self.assertEqual(service.list_duplicates(), [])

        # Висит в отчётах — тишина.
        day4 = self._import(["111-A-1", "111-NEW-1"] + OTHERS)
        self.assertEqual(self._shipment("111-A-1").assignment_status,
                         AssignmentStatus.DELIVERED)
        self.assertEqual(self._export(day4), ["111-NEW-1"])

        # Пропала и вернулась — снова дубль.
        self._import(OTHERS)
        self._import(["111-A-1"] + OTHERS)
        self.assertEqual(self._shipment("111-A-1").assignment_status,
                         AssignmentStatus.DUPLICATE)
        self.assertEqual(len(service.list_duplicates()), 1)

    def test_duplicate_without_gap_stays_and_follows_client_point(self):
        self._issue_and_return()
        client = self.session.execute(select(Client)).scalar_one()
        client.fixed_delivery_point = DeliveryPoint.KOLTSEVAYA_16
        self.session.commit()
        self._import(["111-A-1"] + OTHERS)
        a = self._shipment("111-A-1")
        self.assertEqual(a.assignment_status, AssignmentStatus.DUPLICATE)
        self.assertEqual(a.assigned_point, DeliveryPoint.KOLTSEVAYA_16)
        self.assertEqual(len(DuplicateService(self.session).list_duplicates()), 1)

    def test_duplicate_absent_from_latest_report_is_not_listed(self):
        self._issue_and_return()
        self._import(OTHERS)
        self.assertEqual(DuplicateService(self.session).list_duplicates(), [])

    def test_gap_only_in_incomplete_report_is_not_duplicate(self):
        day1 = self._import(["111-A-1"] + OTHERS)
        self._export(day1)
        self._import(["111-NEW-1"])           # неполный: выпали все 10
        day3 = self._import(["111-A-1", "111-NEW-1"] + OTHERS)

        for posting in ["111-A-1"] + OTHERS:
            shipment = self._shipment(posting)
            self.assertEqual(shipment.assignment_status, AssignmentStatus.TO_SHIP)
            self.assertEqual(shipment.exported_import_session_id, day1.id)
        self.assertEqual(self.importer.skipped_duplicates, 10)
        self.assertEqual(
            json.loads(day3.log_json),
            ["WARNING: Previous import looks incomplete: 10 of 11 shipments missing. "
             "Not treated as duplicates: 10."],
        )
        self.assertEqual(DuplicateService(self.session).list_duplicates(), [])

    def test_longer_gap_is_duplicate_even_after_incomplete_report(self):
        day1 = self._import(["111-A-1"] + OTHERS)
        self._export(day1)
        self._import(OTHERS)                  # A выдана
        self._import(["111-NEW-1"])           # неполный
        self._import(["111-A-1", "111-NEW-1"] + OTHERS)

        self.assertEqual(self._shipment("111-A-1").assignment_status,
                         AssignmentStatus.DUPLICATE)
        self.assertEqual(self.importer.skipped_duplicates, len(OTHERS))

    def test_not_ours_rows_count_towards_incomplete_check(self):
        strangers = [f"222-X-{i}" for i in range(1, 6)]
        day1 = self._import(["111-A-1"] + strangers)
        self._export(day1)
        self._import(["111-NEW-1"])           # выпали A и 5 чужих
        self._import(["111-A-1", "111-NEW-1"] + strangers)
        self.assertEqual(self._shipment("111-A-1").assignment_status, AssignmentStatus.TO_SHIP)
        self.assertEqual(self.importer.skipped_duplicates, 1)

    def test_stale_list_skips_duplicates_and_dismissed(self):
        self._issue_and_return()
        stale = StaleShipmentService(self.session)
        self.assertNotIn("111-A-1", [r.posting_number for r in stale.list_stale()])
        DuplicateService(self.session).dismiss(self._shipment("111-A-1").id)
        self.assertNotIn("111-A-1", [r.posting_number for r in stale.list_stale()])


class IncompleteThresholdTest(unittest.TestCase):
    def test_thresholds(self):
        self.assertFalse(is_incomplete(4, 0))
        self.assertTrue(is_incomplete(5, 95))
        self.assertFalse(is_incomplete(5, 96))
        self.assertTrue(is_incomplete(50, 950))
        self.assertFalse(is_incomplete(0, 0))


if __name__ == "__main__":
    unittest.main()
