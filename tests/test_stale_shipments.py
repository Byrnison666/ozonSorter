"""Залежавшиеся посылки: StaleShipmentService на реальной SQLite и реальных xlsx.

Проверяем выборку (только посылки наших активных клиентов из последнего отчёта),
расчёт дней до даты последнего отчёта, фильтры, сортировку, строковое
представление и Excel-экспорт с подсветкой. Прошедшие дни моделируются правкой
Shipment.first_seen_at / ImportSession.started_at прямо в БД.
"""
import os
import tempfile
import unittest
from datetime import datetime, timedelta

import openpyxl
from sqlalchemy import select

from src.database import DatabaseManager
from src.models import (
    Client, Shipment, ImportSession, DeliveryPoint, AssignmentStatus,
)
from src.services import ImportService
from src.stale_service import (
    StaleShipmentService, StaleRow, COLUMNS, POINT_LABELS,
)

READY = "Готово к выдаче"
RETURN = "Отправить на склад"

# Полдень — чтобы сдвиги на целые дни не упирались в границу суток.
REPORT_TS = datetime(2026, 3, 15, 12, 0, 0)


def _make_report(rows):
    """xlsx отчёта склада. rows: (номер, статус, ячейка) или (номер, статус, ячейка, метка)."""
    path = tempfile.mktemp(suffix=".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Этикетка\nНазвание", "Номер отправления", "Тип", "Статус", "Ячейка"])
    for r in rows:
        posting, status, cell = r[:3]
        label = r[3] if len(r) > 3 else f"LBL-{posting}"
        ws.append([f"{label}\nТовар {posting}", posting, "Обычный", status, cell])
    wb.save(path)
    return path


class StaleShipmentsTest(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        self.db = DatabaseManager(db_path=self.db_path)
        self.db.create_tables()
        self.session = self.db.get_session()
        self.session.add_all([
            Client(ozon_client_id="111", full_name="Иванов", phone="+7900",
                   fixed_delivery_point=DeliveryPoint.KOMSOMOLSKAYA_4),
            Client(ozon_client_id="222", full_name="Петров", phone="+7911",
                   fixed_delivery_point=DeliveryPoint.KOLTSEVAYA_16),
        ])
        self.session.commit()
        self.importer = ImportService(self.session)
        self.service = StaleShipmentService(self.session)
        self._files = []

    def tearDown(self):
        self.session.close()
        os.remove(self.db_path)
        for f in self._files:
            if os.path.exists(f):
                os.remove(f)

    # --- хелперы ---

    def _tmp_xlsx(self):
        path = tempfile.mktemp(suffix=".xlsx")
        self._files.append(path)
        return path

    def _import(self, rows):
        path = _make_report(rows)
        self._files.append(path)
        return self.importer.process_import(path)

    def _import_at(self, rows, started_at=REPORT_TS):
        """Импорт + фиксация даты отчёта, чтобы тест не зависел от now()."""
        imp = self._import(rows)
        self._set_report_date(imp, started_at)
        return imp

    def _set_report_date(self, imp, started_at):
        imp.started_at = started_at
        self.session.commit()

    def _shipment(self, posting):
        return self.session.execute(
            select(Shipment).where(Shipment.posting_number == posting)
        ).scalar_one()

    def _age(self, posting, days, ref=REPORT_TS):
        """Сделать посылку «впервые увиденной» days дней назад относительно ref."""
        ship = self._shipment(posting)
        ship.first_seen_at = ref - timedelta(days=days)
        self.session.commit()

    def _postings(self, rows):
        return [r.posting_number for r in rows]

    # --- 1. нет импортов ---

    def test_no_imports_latest_is_none_and_list_empty(self):
        self.assertIsNone(self.service.latest_import())
        self.assertEqual(self.service.list_stale(), [])

    # --- 2. расчёт дней ---

    def test_days_counted_from_first_seen_to_report_date(self):
        self._import_at([("111-A-1", READY, "A-01")])
        self._age("111-A-1", 5)
        rows = self.service.list_stale()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].days, 5)

    def test_days_zero_when_first_seen_on_report_date(self):
        self._import_at([("111-A-1", READY, "A-01")])
        self._age("111-A-1", 0)
        self.assertEqual(self.service.list_stale()[0].days, 0)

    def test_days_clamped_to_zero_when_first_seen_after_report_date(self):
        self._import_at([("111-A-1", READY, "A-01")])
        self._age("111-A-1", -3)
        self.assertEqual(self.service.list_stale()[0].days, 0)

    # --- 3. дни относительно даты отчёта, а не now ---

    def test_days_relative_to_last_report_date_not_now(self):
        imp = self._import_at([("111-A-1", READY, "A-01")])
        old_report = datetime.now() - timedelta(days=100)
        self._set_report_date(imp, old_report)
        self._age("111-A-1", 4, ref=old_report)
        # От now() было бы ~104 дня.
        self.assertEqual(self.service.list_stale()[0].days, 4)

    def test_reimport_does_not_reset_days(self):
        # Главный инвариант: повторная встреча посылки не сбрасывает первую дату.
        self._import_at([("111-A-1", READY, "A-01")], REPORT_TS - timedelta(days=6))
        self._age("111-A-1", 6)
        self._import_at([("111-A-1", READY, "A-01")])
        self.assertEqual(self.service.list_stale()[0].days, 6)

    def test_same_posting_with_two_barcodes_gives_two_rows(self):
        self._import_at([
            ("111-A-1", READY, "A-01", "ii100"),
            ("111-A-1", READY, "A-01", "ii200"),
        ])
        self.assertEqual(
            sorted(r.product_label for r in self.service.list_stale()),
            ["ii100", "ii200"])

    # --- 4. посылка исчезла из последнего отчёта ---

    def test_shipment_absent_from_latest_report_not_listed(self):
        self._import_at([("111-X-1", READY, "A-01"), ("111-Y-1", READY, "A-02")],
                        REPORT_TS - timedelta(days=1))
        self._import_at([("111-Y-1", READY, "A-02")])
        self.assertEqual(self._postings(self.service.list_stale()), ["111-Y-1"])

    # --- 5. не наши и КТЯ ---

    def test_not_ours_and_kty_excluded(self):
        self._import_at([
            ("111-A-1", READY, "A-01"),
            ("999-B-1", READY, "A-02"),   # клиента 999 в базе нет
            ("ABC-C-1", READY, "A-03"),   # нет цифрового префикса -> КТЯ
        ])
        self.assertEqual(
            self._shipment("999-B-1").assignment_status,
            AssignmentStatus.EXCLUDED_NOT_OURS)
        self.assertEqual(
            self._shipment("ABC-C-1").assignment_status,
            AssignmentStatus.EXCLUDED_KTY)
        self.assertEqual(self._postings(self.service.list_stale()), ["111-A-1"])

    # --- 6. неактивный клиент ---

    def test_inactive_client_shipments_excluded(self):
        self._import_at([("111-A-1", READY, "A-01"), ("222-B-1", READY, "A-02")])
        client = self.session.execute(
            select(Client).where(Client.ozon_client_id == "222")).scalar_one()
        client.is_active = False
        self.session.commit()
        self.assertEqual(self._postings(self.service.list_stale()), ["111-A-1"])

    # --- 7. возвраты ---

    def test_returned_shipment_included_with_return_flag_and_status_cell(self):
        self._import_at([("111-A-1", READY, "A-01"), ("111-R-1", RETURN, "A-02")])
        rows = {r.posting_number: r for r in self.service.list_stale()}
        self.assertEqual(set(rows), {"111-A-1", "111-R-1"})
        self.assertIs(rows["111-R-1"].is_returned, True)
        self.assertEqual(rows["111-R-1"].as_cells()[-1], "Возврат")
        self.assertIs(rows["111-A-1"].is_returned, False)
        self.assertEqual(rows["111-A-1"].as_cells()[-1], "Готово к выдаче")

    # --- 8. фильтры ---

    def test_point_filter(self):
        self._import_at([("111-A-1", READY, "A-01"), ("222-B-1", READY, "A-02")])
        kom = self.service.list_stale(point=DeliveryPoint.KOMSOMOLSKAYA_4)
        kol = self.service.list_stale(point=DeliveryPoint.KOLTSEVAYA_16)
        self.assertEqual(self._postings(kom), ["111-A-1"])
        self.assertEqual(self._postings(kol), ["222-B-1"])
        self.assertEqual(len(self.service.list_stale()), 2)

    def test_min_days_filter_boundary_inclusive(self):
        self._import_at([
            ("111-A-1", READY, "A-01"),
            ("111-B-1", READY, "A-02"),
            ("111-C-1", READY, "A-03"),
        ])
        self._age("111-A-1", 2)
        self._age("111-B-1", 3)
        self._age("111-C-1", 4)
        got = self.service.list_stale(min_days=3)
        self.assertEqual(self._postings(got), ["111-C-1", "111-B-1"])
        self.assertEqual(len(self.service.list_stale(min_days=0)), 3)
        self.assertEqual(self.service.list_stale(min_days=5), [])

    # --- 9. сортировка ---

    def test_sorted_by_days_desc_then_cell_then_posting(self):
        self._import_at([
            ("111-P1-1", READY, "B-02"),
            ("111-P2-1", READY, "A-01"),
            ("111-P3-1", READY, "A-01"),
            ("111-P4-1", READY, "C-03"),
        ])
        self._age("111-P1-1", 2)
        self._age("111-P2-1", 2)
        self._age("111-P3-1", 2)
        self._age("111-P4-1", 9)
        got = self.service.list_stale()
        self.assertEqual(
            self._postings(got),
            ["111-P4-1", "111-P2-1", "111-P3-1", "111-P1-1"])

    def test_cells_sorted_naturally_not_as_text(self):
        # Как текст «10-1» шло бы раньше «2-1» — по стеллажам так не ходят.
        self._import_at([
            ("111-P1-1", READY, "10-1"),
            ("111-P2-1", READY, "2-10"),
            ("111-P3-1", READY, "2-3"),
        ])
        self.assertEqual(
            [r.cell for r in self.service.list_stale()], ["2-3", "2-10", "10-1"])

    # --- 10. as_cells ---

    def _sample_row(self, **kw):
        base = dict(
            days=4, cell="A-01", posting_number="111-A-1", product_label="LBL",
            product_name="Товар", client_name="Иванов", ozon_client_id="111",
            phone="+7900", point=DeliveryPoint.KOLTSEVAYA_16,
            first_seen_at=datetime(2026, 3, 5, 9, 30), is_returned=False,
        )
        base.update(kw)
        return StaleRow(**base)

    def test_as_cells_order_length_date_format_and_point_label(self):
        cells = self._sample_row().as_cells()
        self.assertEqual(len(cells), len(COLUMNS))
        self.assertEqual(cells, [
            4, "A-01", "111-A-1", "LBL", "Товар", "Иванов", "111", "+7900",
            POINT_LABELS[DeliveryPoint.KOLTSEVAYA_16], "05.03.2026",
            "Готово к выдаче",
        ])

    def test_as_cells_from_service_has_empty_strings_for_missing_client_fields(self):
        self.session.add(Client(
            ozon_client_id="333", full_name=None, phone=None,
            fixed_delivery_point=DeliveryPoint.KOMSOMOLSKAYA_4))
        self.session.commit()
        self._import_at([("333-A-1", READY, "A-01")])
        row = self.service.list_stale()[0]
        cells = row.as_cells()
        self.assertEqual(cells[COLUMNS.index("Клиент")], "")
        self.assertEqual(cells[COLUMNS.index("Телефон")], "")
        self.assertEqual(cells[COLUMNS.index("Ozon ID")], "333")

    # --- 11. экспорт ---

    def test_export_returns_row_count_and_writes_header_and_values(self):
        self._import_at([("111-A-1", READY, "A-01"), ("222-B-1", RETURN, "A-02")])
        self._age("111-A-1", 1)
        self._age("222-B-1", 2)
        rows = self.service.list_stale()
        out = self._tmp_xlsx()
        self.assertEqual(self.service.export_xlsx(rows, out), 2)

        ws = openpyxl.load_workbook(out).active
        self.assertEqual(ws.title, "Залежавшиеся")
        header = [c.value for c in ws[1]]
        self.assertEqual(header, COLUMNS)
        self.assertTrue(all(c.font.bold for c in ws[1]))
        self.assertEqual(
            [c.value for c in ws[2]], rows[0].as_cells())
        self.assertEqual(
            [c.value for c in ws[3]], rows[1].as_cells())
        self.assertEqual(ws[2][-1].value, "Возврат")  # days=2 у возврата идёт первым

    def test_export_fill_by_thresholds(self):
        self._import_at([
            ("111-D-1", READY, "A-01"),
            ("111-W-1", READY, "A-02"),
            ("111-N-1", READY, "A-03"),
        ])
        self._age("111-D-1", 7)
        self._age("111-W-1", 3)
        self._age("111-N-1", 2)
        rows = self.service.list_stale()
        out = self._tmp_xlsx()
        self.service.export_xlsx(rows, out)
        ws = openpyxl.load_workbook(out).active

        by_posting = {}
        for excel_row in ws.iter_rows(min_row=2):
            by_posting[excel_row[2].value] = excel_row

        for c in by_posting["111-D-1"]:
            self.assertEqual(c.fill.fill_type, "solid")
            self.assertTrue(c.fill.start_color.rgb.endswith("FFE5E5"),
                            c.fill.start_color.rgb)
        for c in by_posting["111-W-1"]:
            self.assertEqual(c.fill.fill_type, "solid")
            self.assertTrue(c.fill.start_color.rgb.endswith("FFF4D6"),
                            c.fill.start_color.rgb)
        for c in by_posting["111-N-1"]:
            self.assertIsNone(c.fill.fill_type)

    def test_export_freeze_panes(self):
        out = self._tmp_xlsx()
        self.service.export_xlsx([], out)
        self.assertEqual(openpyxl.load_workbook(out).active.freeze_panes, "A2")

    def test_export_empty_list_writes_only_header(self):
        out = self._tmp_xlsx()
        self.assertEqual(self.service.export_xlsx([], out), 0)
        ws = openpyxl.load_workbook(out).active
        self.assertEqual(ws.max_row, 1)
        self.assertEqual([c.value for c in ws[1]], COLUMNS)

    # --- 12. обновление ячейки при повторной встрече ---

    def test_reimport_updates_cell_and_list_shows_new_cell(self):
        self._import_at([("111-A-1", READY, "A-01")], REPORT_TS - timedelta(days=1))
        self._import_at([("111-A-1", READY, "B-07")])
        self.assertEqual(self._shipment("111-A-1").cell, "B-07")
        self.assertEqual(self.service.list_stale()[0].cell, "B-07")

    def test_reimport_with_empty_cell_keeps_known_cell(self):
        # Фактическое поведение: строка с пустой ячейкой не пропускается
        # (посылка остаётся в последнем отчёте), но известную ячейку не затирает.
        self._import_at([("111-A-1", READY, "A-01")], REPORT_TS - timedelta(days=1))
        imp2 = self._import_at([("111-A-1", READY, None)])
        ship = self._shipment("111-A-1")
        self.assertEqual(ship.cell, "A-01")
        self.assertEqual(ship.last_seen_import_session_id, imp2.id)
        rows = self.service.list_stale()
        self.assertEqual(self._postings(rows), ["111-A-1"])
        self.assertEqual(rows[0].cell, "A-01")

    def test_reimport_with_whitespace_cell_keeps_known_cell(self):
        self._import_at([("111-A-1", READY, "A-01")], REPORT_TS - timedelta(days=1))
        self._import_at([("111-A-1", READY, "   ")])
        self.assertEqual(self._shipment("111-A-1").cell, "A-01")

    # --- 13. последний отчёт = max id ---

    def test_latest_import_is_max_id_after_reimport_of_earlier_file(self):
        path1 = _make_report([("111-A-1", READY, "A-01")])
        path2 = _make_report([("111-A-1", READY, "A-01"), ("111-B-1", READY, "A-02")])
        self._files += [path1, path2]
        self.importer.process_import(path1)
        self.importer.process_import(path2)
        self.assertEqual(
            sorted(self._postings(self.service.list_stale())),
            ["111-A-1", "111-B-1"])

        # Повторный импорт более раннего файла: он теперь «последний» по id.
        imp3 = self.importer.process_import(path1)
        self.assertEqual(self.service.latest_import().id, imp3.id)
        self.assertEqual(self._postings(self.service.list_stale()), ["111-A-1"])


if __name__ == "__main__":
    unittest.main()
