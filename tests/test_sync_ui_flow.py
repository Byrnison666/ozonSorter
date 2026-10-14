"""Сквозные сценарии синхронизации через интерфейс (Qt offscreen): главное окно,
экран «Настройки», экраны данных и локальный «Яндекс.Диск» (REST API).

Настоящие модальные диалоги в тестах не показываются (они заблокировали бы
прогон): вопросы пользователю и итоговые сообщения подменены, а тест задаёт,
что «нажал» пользователь, и проверяет, какие сообщения он увидел.
"""
import gzip
import json
import os
import shutil
import socket
import tempfile
import threading
import time
import unittest
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import openpyxl
from PySide2.QtCore import Qt, QTimer
from PySide2.QtGui import QCloseEvent
from PySide2.QtWidgets import QApplication, QDialog, QMessageBox
from sqlalchemy import select

from src.database import DatabaseManager
from src.models import Client, DeliveryPoint
from src.sync_config import SyncConfig, load_config, save_config
from src.sync_service import (
    CancelToken, RemoteDataError, RemoteMeta, SyncCancelled, SyncCheck, SyncError, SyncService,
    SyncStatus,
)
from src.ui import clients_screen, sync_flow
from src.ui.main_window import MainWindow
from src.yadisk_api import YaDiskClient, YaDiskError
from yadisk_fake import FakeYaDiskServer

# Настоящие функции: в тестах они подменяются (ни настоящего интернета, ни
# настоящих диалогов), но сами тоже проверяются.
REAL_INTERNET_REACHABLE = sync_flow.internet_reachable
REAL_ASK_RETRY_OFFLINE = sync_flow._ask_retry_offline
REAL_ASK_CONFLICT = sync_flow._ask_conflict
REAL_ASK_PULL = sync_flow._ask_pull
REAL_ASK_EMPTY_DISK = sync_flow._ask_push_to_empty_disk

CLEAN = "невыложенных изменений нет"
DIRTY = "есть невыложенные изменения"
UNAVAILABLE = "Яндекс.Диск недоступен"


def _add_client(db_path, ozon_id):
    db = DatabaseManager(db_path=db_path)
    db.create_tables()
    session = db.get_session()
    session.add(Client(ozon_client_id=ozon_id, full_name=f"Клиент {ozon_id}",
                       fixed_delivery_point=DeliveryPoint.KOMSOMOLSKAYA_4))
    session.commit()
    session.close()
    db.engine.dispose()


def _client_ids(db_path):
    db = DatabaseManager(db_path=db_path)
    session = db.get_session()
    try:
        return sorted(session.execute(select(Client.ozon_client_id)).scalars().all())
    finally:
        session.close()
        db.engine.dispose()


class FakeBox:
    """Подмена QMessageBox из message_box(): «нажимает» кнопку с заданной подписью."""
    choice = None

    def __init__(self, *args, **kwargs):
        self.buttons = {}

    def addButton(self, text, role):
        button = object()
        self.buttons[text] = button
        return button

    def setDefaultButton(self, button):
        pass

    def setEscapeButton(self, button):
        pass

    def exec_(self):
        pass

    def clickedButton(self):
        return self.buttons.get(FakeBox.choice)


class SyncUiTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.server = FakeYaDiskServer()
        self.dir = tempfile.mkdtemp(prefix="ozon-ui-")
        self.db_path = os.path.join(self.dir, "ozon_sorter.db")
        self.messages = []          # (заголовок, текст) показанных сообщений
        self.other_dir = tempfile.mkdtemp(prefix="ozon-ui-other-")
        self.other_db = os.path.join(self.other_dir, "ozon_sorter.db")

        self._patch_value("src.sync_config.CONFIG_PATH", os.path.join(self.dir, "sync.json"))
        self._patch_value("src.sync_config.API_URL", self.server.api_url)
        self._patch_value("src.sync_config.REMOTE_DIR", "OzonSorter")
        self._patch_value("src.ui.sync_flow.STATE_PATH", os.path.join(self.dir, "sync_state.json"))
        self._patch_value("src.sync_config.STATE_PATH", os.path.join(self.dir, "sync_state.json"))
        record = lambda parent, icon, title, text: self.messages.append((title, text))
        self._patch("src.ui.sync_flow.show_message", side_effect=record)
        self._patch("src.ui.main_window.show_message", side_effect=record)
        self._patch("src.ui.settings_screen.show_message", side_effect=record)
        self.ask_conflict = self._patch("src.ui.sync_flow._ask_conflict", return_value="")
        self.ask_pull = self._patch("src.ui.sync_flow._ask_pull", return_value=True)
        self.ask_empty = self._patch("src.ui.sync_flow._ask_push_to_empty_disk", return_value=True)
        self.ask_retry = self._patch("src.ui.sync_flow._ask_retry_offline", return_value=False)
        self.internet = self._patch("src.ui.sync_flow.internet_reachable", return_value=True)
        self.ask_yes_no = self._patch("src.ui.main_window.ask_yes_no", return_value=False)
        # Повторы запросов при сбое Диска — без настоящих пауз.
        self._patch_value("src.yadisk_api._RETRY_DELAYS", (0, 0))

        self.config = SyncConfig(token=self.server.token, uid=self.server.uid,
                                 login=self.server.login, device_name="ПК")
        save_config(self.config)
        _add_client(self.db_path, "111")
        self.db = DatabaseManager(db_path=self.db_path)
        self.db.create_tables()
        self.window = MainWindow(self.db)
        self.settings = self.window._screens["settings"]

        # Второе устройство — напрямую через сервис, без интерфейса.
        other_db = DatabaseManager(db_path=self.other_db)
        other_db.create_tables()
        other_db.engine.dispose()       # на Windows открытое соединение держит файл
        self.other = SyncService(
            self.other_db, YaDiskClient(self.server.token, self.server.api_url, timeout=5),
            "OzonSorter", os.path.join(self.other_dir, "sync_state.json"), "Телефон",
            target_id=self.config.target_id,
        )

    def tearDown(self):
        sync_flow.stop_pending_work()
        self.window.release_db()
        self.window.deleteLater()
        self.app.processEvents()
        self.server.stop()
        shutil.rmtree(self.dir, ignore_errors=True)
        shutil.rmtree(self.other_dir, ignore_errors=True)

    # --- помощники ---

    def _patch(self, target, **kwargs):
        """mock.patch с автоматической отменой; возвращает подмену (Mock)."""
        patcher = mock.patch(target, **kwargs)
        started = patcher.start()
        self.addCleanup(patcher.stop)
        return started

    def _patch_value(self, target, value):
        patcher = mock.patch(target, value)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _wait(self, condition, timeout=10.0):
        deadline = time.monotonic() + timeout
        while not condition():
            self.app.processEvents()
            time.sleep(0.01)
            self.assertLess(time.monotonic(), deadline, "ожидание не дождалось условия")

    def _push_from_other(self, ozon_id):
        """Другое устройство берёт текущую базу с Диска (если она есть), добавляет
        клиента и выкладывает новую версию."""
        check = self.other.check()
        if check.status == SyncStatus.REMOTE_AHEAD:
            self.other.pull(check.meta)
        _add_client(self.other_db, ozon_id)
        self.other.push(self.other.check().meta)

    def _indicator(self):
        return self.window.sync_label.text()

    def _texts(self):
        return [text for _title, text in self.messages]

    def _clients_on_screen(self):
        table = self.window._screens["clients"].table
        return sorted(table.item(r, 0).text() for r in range(table.rowCount()))

    def _close(self, choice=None):
        FakeBox.choice = choice
        event = QCloseEvent()
        with mock.patch("src.ui.main_window.message_box", side_effect=FakeBox):
            self.window.closeEvent(event)
        return event.isAccepted()


class SettingsAndManualSyncTest(SyncUiTestCase):
    def test_settings_show_logged_in_account(self):
        self.assertIn("tester", self.settings.account_label.text())
        self.assertTrue(self.settings.logout_btn.isVisibleTo(self.settings))
        self.assertFalse(self.settings.login_btn.isVisibleTo(self.settings))
        self.assertEqual(self.settings.device_edit.text(), "ПК")

    def test_sync_button_pushes_local_base_and_announces_success(self):
        self.settings.on_sync()
        self.assertIn("/OzonSorter/meta.json", self.server.files)
        self.assertTrue(any("выложены" in t for t in self._texts()))
        self.assertIn("Синхронизировано", self.settings.status_label.text())
        self.assertEqual(self.other.check().status, SyncStatus.REMOTE_AHEAD)

    def test_up_to_date_is_announced_on_manual_sync(self):
        self.settings.on_sync()
        self.messages.clear()
        self.settings.on_sync()
        self.assertTrue(any("совпадает" in t for t in self._texts()))

    def test_sync_button_pulls_newer_base_without_extra_question(self):
        self.settings.on_sync()
        self.messages.clear()
        self._push_from_other("222")
        old_clients_screen = self.window._screens["clients"]
        self.settings.on_sync()
        self.ask_pull.assert_not_called()             # нажатие кнопки уже было согласием
        self.assertTrue(any("загружена" in t for t in self._texts()))
        self.assertIsNot(self.window._screens["clients"], old_clients_screen)
        self.assertEqual(self._clients_on_screen(), ["111", "222"])
        self.assertIs(self.window._screens["settings"], self.settings)

    def test_database_usable_after_push(self):
        self.settings.on_sync()
        session = self.db.get_session()
        try:
            self.assertEqual(session.execute(select(Client.ozon_client_id)).scalars().all(), ["111"])
        finally:
            session.close()

    def test_revoked_token_shows_readable_error(self):
        self.server.token = "other-token"       # доступ отозван на стороне Яндекса
        self.settings.on_sync()
        self.assertTrue(any("не принял вход" in t for t in self._texts()))
        self.assertIn("не выполнена", self.settings.status_label.text())

    def test_sync_without_login_asks_to_log_in(self):
        save_config(SyncConfig(device_name="ПК"))
        with mock.patch.object(QMessageBox, "information") as info:
            self.settings.on_sync()
        info.assert_called_once()
        self.assertEqual(self.messages, [])

    def test_manual_button_in_sidebar_announces(self):
        self.window.sync_btn.click()
        self.assertTrue(any("выложены" in t for t in self._texts()))


class ConflictAndSafetyTest(SyncUiTestCase):
    def _make_conflict(self):
        self.settings.on_sync()
        self.messages.clear()
        self._push_from_other("222")
        self.window.release_db()
        _add_client(self.db_path, "333")

    def test_conflict_take_remote(self):
        self._make_conflict()
        self.ask_conflict.return_value = "remote"
        self.settings.on_sync()
        check = self.ask_conflict.call_args[0][1]
        self.assertEqual(check.meta.device, "Телефон")
        self.assertEqual(self._clients_on_screen(), ["111", "222"])

    def test_conflict_keep_local(self):
        self._make_conflict()
        self.ask_conflict.return_value = "local"
        self.settings.on_sync()
        self.other.pull(self.other.check().meta)
        self.assertEqual(_client_ids(self.other_db), ["111", "333"])

    def test_conflict_cancel_changes_nothing(self):
        self._make_conflict()
        files_before = dict(self.server.files)
        self.assertFalse(self.window.sync_now(announce=True))
        self.assertEqual(self.server.files, files_before)
        self.assertEqual(_client_ids(self.db_path), ["111", "333"])
        self.assertEqual(self.messages, [])            # об успехе не сообщаем

    def test_remote_disk_emptied_asks_before_pushing_again(self):
        self.settings.on_sync()
        self.server.files.clear()
        self.server.dirs = {"/"}
        self.ask_empty.return_value = False
        self.assertFalse(self.window.sync_now())
        self.ask_empty.assert_called_once()
        self.assertNotIn("/OzonSorter/meta.json", self.server.files)
        self.ask_empty.return_value = True
        self.assertTrue(self.window.sync_now())
        self.assertIn("/OzonSorter/meta.json", self.server.files)

    def test_older_disk_version_asks_even_on_manual_sync(self):
        self.settings.on_sync()
        meta_rev1 = self.server.files["/OzonSorter/meta.json"]
        self._push_from_other("222")
        self.window.sync_now()                          # A на rev2
        self.server.files["/OzonSorter/meta.json"] = meta_rev1
        self.ask_pull.return_value = False
        self.assertFalse(self.window.sync_now())
        self.ask_pull.assert_called_once()
        self.assertTrue(self.ask_pull.call_args[0][1].remote_is_older)
        self.assertEqual(_client_ids(self.db_path), ["111", "222"])

    def test_failed_reload_restores_previous_base(self):
        self.settings.on_sync()
        self.messages.clear()
        self._push_from_other("222")
        real_reload = self.window.reload_db
        calls = {"n": 0}

        def reload_failing_once():
            calls["n"] += 1
            if calls["n"] == 1:
                raise RuntimeError("не открылась")
            real_reload()

        with mock.patch.object(self.window, "reload_db", side_effect=reload_failing_once):
            self.assertFalse(self.window.sync_now(announce=True))
        self.assertTrue(any("возвращена прежняя" in t for t in self._texts()))
        self.assertEqual(_client_ids(self.db_path), ["111"])
        self.assertEqual(calls["n"], 2)

    def test_message_box_shows_text_as_is(self):
        box = sync_flow.message_box(None, QMessageBox.Information, "t", '<img src="file://x/y.png">')
        self.assertEqual(box.textFormat(), Qt.PlainText)


class OfflineTest(SyncUiTestCase):
    def _go_online_and_retry(self, *args):
        self.server.offline = False
        return True

    def test_offline_change_offers_retry_then_confirms_success(self):
        self.server.offline = True
        self.ask_retry.side_effect = self._go_online_and_retry
        self.assertTrue(self.window.sync_now())
        self.ask_retry.assert_called_once()
        self.assertTrue(self.ask_retry.call_args[0][1])          # сказано, что есть невыложенные правки
        self.assertTrue(any("выложены" in t for t in self._texts()))
        self.assertIn(CLEAN, self._indicator())
        self.assertIn("/OzonSorter/meta.json", self.server.files)

    def test_offline_and_user_chooses_later_keeps_changes_marked(self):
        self.server.offline = True
        self.ask_retry.return_value = False
        self.assertFalse(self.window.sync_now())
        self.assertEqual(self.messages, [])
        self.assertIn(DIRTY, self._indicator())

    def test_retry_that_still_fails_asks_again(self):
        self.server.offline = True
        answers = iter([True, True, False])
        self.ask_retry.side_effect = lambda *a: next(answers)
        self.assertFalse(self.window.sync_now())
        self.assertEqual(self.ask_retry.call_count, 3)
        self.assertEqual(self.messages, [])

    def test_data_change_syncs_automatically_and_stays_quiet_when_online(self):
        self.window.on_data_changed()
        self.assertIn("/OzonSorter/meta.json", self.server.files)
        self.assertEqual(self.messages, [])                      # успех без сбоев — без лишних окон
        self.assertIn(CLEAN, self._indicator())

    def _change_while_disk_is_down(self):
        """Синхронизированы, затем правка, а Диск в это время не работает."""
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.window.on_data_changed()

    def _assert_working_locally(self):
        self.ask_retry.assert_not_called()
        self.assertEqual(self.messages, [])
        self.assertIn(UNAVAILABLE, self._indicator())
        self.assertIn(DIRTY, self._indicator())
        self.assertIn("Последняя синхронизация", self._indicator())
        self.assertEqual(_client_ids(self.db_path), ["111", "222"])

    def test_data_change_offline_is_silent_and_marks_indicator(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.offline = True
        self.window.on_data_changed()
        self._assert_working_locally()

    def test_data_change_with_server_error_is_silent(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.server_error = 500
        self.window.on_data_changed()
        self._assert_working_locally()

    def test_failed_disk_operation_on_push_is_silent(self):
        # Ошибка 500 могла быть и сбоем переноса файла на самом Диске.
        self._patch_value("src.yadisk_api._DIRECT_UPLOAD_MAX", 0)
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.fail_ops = True
        self.window.on_data_changed()
        self._assert_working_locally()
        self.assertEqual(self.other.check().status, SyncStatus.REMOTE_AHEAD)
        self.other.pull(self.other.check().meta)
        self.assertEqual(_client_ids(self.other_db), ["111"])    # на Диске прежняя целая база

    def test_no_new_attempts_during_pause_after_failure(self):
        self.window.on_data_changed()
        self.server.server_error = 503
        _add_client(self.db_path, "222")
        self.window.on_data_changed()
        requests = len(self.server.requests)
        for ozon_id in ("333", "444"):
            _add_client(self.db_path, ozon_id)
            self.window.on_data_changed()
        self.assertEqual(len(self.server.requests), requests)   # не ждём Диск на каждой правке
        self.assertIn(UNAVAILABLE, self._indicator())

    def test_after_pause_change_is_pushed_and_mark_cleared(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.server_error = 500
        self.window.on_data_changed()
        self.server.server_error = None
        self.window._auto_sync_after = 0.0          # пауза истекла
        _add_client(self.db_path, "333")
        self.window.on_data_changed()
        self.assertEqual(self.messages, [])
        self.assertIn(CLEAN, self._indicator())
        self.assertNotIn(UNAVAILABLE, self._indicator())
        self._push_from_other("444")
        self.assertEqual(_client_ids(self.other_db), ["111", "222", "333", "444"])

    def test_manual_sync_ignores_pause_and_explains_local_work(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.server_error = 500
        self.window.on_data_changed()
        requests = len(self.server.requests)
        self.assertFalse(self.window.sync_now(announce=True))
        self.assertGreater(len(self.server.requests), requests)
        self.assertTrue(any("работать можно" in t for t in self._texts()))
        self.assertIn(UNAVAILABLE, self._indicator())
        self.server.server_error = None
        self.assertTrue(self.window.sync_now(announce=True))
        self.assertNotIn(UNAVAILABLE, self._indicator())

    def test_insufficient_storage_is_not_hidden(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.fail_next["GET"] = 507
        self.window.on_data_changed()
        self.assertTrue(any("закончилось место" in t for t in self._texts()))
        self.assertNotIn(UNAVAILABLE, self._indicator())

    def test_indicator_reflects_state(self):
        self.assertIn(DIRTY, self._indicator())
        self.window.sync_now()
        self.assertIn(CLEAN, self._indicator())
        _add_client(self.db_path, "222")
        self.window.refresh_sync_indicator()
        self.assertIn(DIRTY, self._indicator())

    def test_indicator_hidden_without_config(self):
        save_config(SyncConfig())
        self.window.refresh_sync_indicator()
        self.assertFalse(self.window.sync_label.isVisibleTo(self.window))
        self.assertFalse(self.window.sync_btn.isVisibleTo(self.window))

    def test_second_sync_while_busy_is_ignored(self):
        self.window._sync_busy = True
        self.assertFalse(self.window.sync_now())
        self.assertNotIn("/OzonSorter/meta.json", self.server.files)


class StartupCheckTest(SyncUiTestCase):
    def _startup(self):
        self.window.startup_sync()
        self._wait(lambda: self.window._startup_check is None)

    def test_not_configured_does_nothing(self):
        save_config(SyncConfig())
        self._startup()
        self.assertEqual(self.server.requests, [])

    def test_unreadable_token_is_reported(self):
        config = SyncConfig(uid=self.server.uid)
        config.token_unreadable = True
        with mock.patch("src.ui.main_window.load_config", return_value=config):
            self._startup()
        self.assertTrue(any("Синхронизация выключена" in title for title, _t in self.messages))
        self.assertEqual(self.server.requests, [])

    def test_offline_without_local_changes_is_silent(self):
        self.window.sync_now()                      # синхронизированы
        self.server.offline = True
        self._startup()
        self.ask_retry.assert_not_called()
        self.assertEqual(self.messages, [])
        self.assertIn(UNAVAILABLE, self._indicator())
        self.assertIn(CLEAN, self._indicator())

    def test_offline_with_unsynced_changes_is_silent_and_marked(self):
        self.server.offline = True
        self._startup()
        self.ask_retry.assert_not_called()
        self.assertEqual(self.messages, [])
        self.assertIn(UNAVAILABLE, self._indicator())
        self.assertIn(DIRTY, self._indicator())

    def test_changes_made_while_disk_was_down_are_pushed_at_next_start(self):
        self.window.sync_now()
        self.messages.clear()
        self.server.server_error = 500
        self._startup()
        _add_client(self.db_path, "222")            # работа с локальной базой
        self.window.on_data_changed()
        self.assertEqual(self.messages, [])
        self.server.server_error = None             # Яндекс починил сервера
        self._startup()                             # следующий запуск программы
        self.assertTrue(any("выложены" in t for t in self._texts()))
        self.assertIn(CLEAN, self._indicator())
        self.assertNotIn(UNAVAILABLE, self._indicator())
        self._push_from_other("333")
        self.assertEqual(_client_ids(self.other_db), ["111", "222", "333"])

    def test_up_to_date_is_silent(self):
        self.window.sync_now()
        self.messages.clear()
        self._startup()
        self.assertEqual(self.messages, [])
        self.ask_pull.assert_not_called()

    def test_newer_base_on_disk_is_offered_and_loaded(self):
        self.window.sync_now()
        self._push_from_other("222")
        self._startup()
        self.ask_pull.assert_called_once()
        self.assertEqual(self._clients_on_screen(), ["111", "222"])

    def test_newer_base_declined_leaves_local_base(self):
        self.window.sync_now()
        self._push_from_other("222")
        self.ask_pull.return_value = False
        self._startup()
        self.assertEqual(_client_ids(self.db_path), ["111"])

    def test_unsynced_local_changes_are_pushed_and_announced(self):
        self._startup()
        self.assertIn("/OzonSorter/meta.json", self.server.files)
        self.assertTrue(any("выложены" in t for t in self._texts()))

    def test_server_error_at_startup_is_silent_and_marked(self):
        self.server.server_error = 500
        self._startup()
        self.assertEqual(self.messages, [])
        self.assertIn(UNAVAILABLE, self._indicator())

    def test_short_server_failure_at_startup_passes_unnoticed(self):
        self.window.sync_now()
        self.messages.clear()
        self.server.fail_next["GET"] = 502          # повтор запроса его переживёт
        self._startup()
        self.assertEqual(self.messages, [])
        self.assertNotIn(UNAVAILABLE, self._indicator())

    def test_revoked_token_at_startup_is_shown(self):
        self.server.token = "other-token"
        self._startup()
        self.assertTrue(any("не принял вход" in t for t in self._texts()))


class GarbageResponseTest(SyncUiTestCase):
    """Сбой в дата-центре Яндекса бывает не только 500: Диск отвечает «200 OK»
    HTML-страницей, обрезанным архивом или файлом с чужой суммой. Тихие попытки
    считают это недоступностью — отметка и пауза, без окон."""

    def _startup(self):
        self.window.startup_sync()
        self._wait(lambda: self.window._startup_check is None)

    def _newer_base_on_disk(self):
        """Синхронизированы, затем другое устройство выложило новую версию."""
        self.window.sync_now()
        self._push_from_other("222")
        self.messages.clear()

    def _replace_current_revision(self, data):
        meta = json.loads(self.server.files["/OzonSorter/meta.json"])
        self.server.files["/OzonSorter/" + meta["file"]] = data

    def _assert_quiet_unavailable(self):
        self.assertEqual(self.messages, [])
        self.assertIn(UNAVAILABLE, self._indicator())
        self.assertGreater(self.window._auto_sync_after, time.monotonic())

    def test_garbage_api_answer_on_change_is_silent_logged_and_paused(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.garbage_api = True
        with self.assertLogs("src.ui.sync_flow", "WARNING"):
            self.window.on_data_changed()
        self._assert_quiet_unavailable()
        self.assertIn(DIRTY, self._indicator())
        self.assertEqual(_client_ids(self.db_path), ["111", "222"])
        requests = len(self.server.requests)
        _add_client(self.db_path, "333")
        self.window.on_data_changed()
        self.assertEqual(len(self.server.requests), requests)   # пауза: Диск не дёргаем

    def test_garbage_meta_at_startup_is_silent_and_marked(self):
        self.window.sync_now()
        self.messages.clear()
        self.server.garbage_gets = True
        with self.assertLogs("src.ui.main_window", "WARNING"):
            self._startup()
        self._assert_quiet_unavailable()
        self.assertIn(CLEAN, self._indicator())

    def test_checksum_mismatch_on_quiet_pull_keeps_local_base(self):
        self._newer_base_on_disk()
        self._replace_current_revision(gzip.compress(b"not the base"))
        self._startup()
        self._assert_quiet_unavailable()
        self.assertEqual(_client_ids(self.db_path), ["111"])

    def test_broken_archive_on_quiet_pull_keeps_local_base(self):
        self._newer_base_on_disk()
        self._replace_current_revision(b"<html>502</html>")
        self._startup()
        self._assert_quiet_unavailable()
        self.assertEqual(_client_ids(self.db_path), ["111"])

    def test_manual_sync_shows_window_and_marks(self):
        self.window.on_data_changed()
        _add_client(self.db_path, "222")
        self.server.garbage_api = True
        self.assertFalse(self.window.sync_now(announce=True))
        self.assertEqual([title for title, _t in self.messages], [UNAVAILABLE])
        self.assertTrue(any("непонятный ответ" in t and "работать можно" in t
                            for t in self._texts()))
        self.assertIn(UNAVAILABLE, self._indicator())
        self.server.garbage_api = False
        self.assertTrue(self.window.sync_now())
        self.assertNotIn(UNAVAILABLE, self._indicator())
        self._push_from_other("333")
        self.assertEqual(_client_ids(self.other_db), ["111", "222", "333"])

    def test_manual_pull_with_bad_checksum_explains_local_work(self):
        self._newer_base_on_disk()
        self._replace_current_revision(gzip.compress(b"not the base"))
        self.assertFalse(self.window.sync_now(announce=True))
        self.assertEqual([title for title, _t in self.messages], [UNAVAILABLE])
        self.assertTrue(any("контрольной суммой" in t and "работать можно" in t
                            for t in self._texts()))
        self.assertEqual(_client_ids(self.db_path), ["111"])

    def test_glitch_classification(self):
        self.assertTrue(sync_flow.is_disk_glitch(YaDiskError("html", 200)))
        self.assertTrue(sync_flow.is_disk_glitch(YaDiskError("down", 503)))
        self.assertTrue(sync_flow.is_disk_glitch(YaDiskError("offline")))
        self.assertTrue(sync_flow.is_disk_glitch(RemoteDataError("broken")))
        self.assertFalse(sync_flow.is_disk_glitch(YaDiskError("auth", 401)))
        self.assertFalse(sync_flow.is_disk_glitch(YaDiskError("full", 507)))
        self.assertFalse(sync_flow.is_disk_glitch(YaDiskError("tls", tls=True)))
        self.assertFalse(sync_flow.is_disk_glitch(SyncError("other")))
        self.assertFalse(YaDiskError("html", 200).is_unavailable)  # отдельный признак


class CloseWindowTest(SyncUiTestCase):
    def test_close_without_changes_needs_no_question(self):
        self.window.sync_now()
        self.assertTrue(self._close(choice=None))

    def test_close_with_unsynced_changes_and_skip(self):
        self.assertTrue(self._close("Закрыть без выгрузки"))
        self.assertNotIn("/OzonSorter/meta.json", self.server.files)

    def test_close_with_unsynced_changes_and_push(self):
        self.assertTrue(self._close("Выложить"))
        self.assertIn("/OzonSorter/meta.json", self.server.files)

    def test_close_cancelled(self):
        self.assertFalse(self._close("Отмена"))
        self.assertNotIn("/OzonSorter/meta.json", self.server.files)

    def test_close_push_fails_offline_and_user_keeps_window_open(self):
        self.server.offline = True
        self.ask_yes_no.return_value = False
        self.assertFalse(self._close("Выложить"))
        self.ask_retry.assert_not_called()
        self.ask_yes_no.assert_called_once()
        text = self.ask_yes_no.call_args[0][2]
        self.assertIn("выложатся при следующем запуске", text)
        self.assertIs(self.ask_yes_no.call_args[1]["default_yes"], True)

    def test_close_push_fails_offline_and_user_closes_anyway(self):
        self.server.offline = True
        self.ask_yes_no.return_value = True
        self.assertTrue(self._close("Выложить"))

    def test_close_push_fails_with_server_error_and_user_closes(self):
        self.server.server_error = 500
        self.ask_yes_no.return_value = True
        self.assertTrue(self._close("Выложить"))
        self.assertEqual(self.messages, [])
        self.assertIn("выложатся при следующем запуске", self.ask_yes_no.call_args[0][2])

    def test_close_push_declined_conflict_asks_again(self):
        self.window.sync_now()
        self._push_from_other("222")
        self.window.release_db()
        _add_client(self.db_path, "333")
        self.ask_conflict.return_value = ""             # диалог конфликта отменён
        self.ask_yes_no.return_value = False
        self.assertFalse(self._close("Выложить"))
        self.ask_yes_no.assert_called_once()
        self.assertIn("так и не выложены", self.ask_yes_no.call_args[0][2])
        self.assertIs(self.ask_yes_no.call_args[1]["default_yes"], False)

    def test_close_ignored_while_sync_is_running(self):
        self.window._sync_busy = True
        self.assertFalse(self._close("Закрыть без выгрузки"))

    def test_failure_in_close_check_does_not_lock_the_user_in(self):
        with mock.patch("src.ui.main_window.load_config", side_effect=RuntimeError("boom")):
            event = QCloseEvent()
            self.window.closeEvent(event)
        self.assertTrue(event.isAccepted())


class ProgressDialogTest(SyncUiTestCase):
    def _reject_modal_after(self, ms):
        def reject():
            dialog = QApplication.activeModalWidget()
            if dialog is not None:
                dialog.reject()
        QTimer.singleShot(ms, reject)

    def test_cancel_before_commit_abandons_the_operation(self):
        gate = threading.Event()
        self._reject_modal_after(100)
        with self.assertRaises(SyncCancelled):
            sync_flow.run_with_progress(self.window, "t", lambda token: gate.wait(10) and "x")
        self.assertTrue(sync_flow._abandoned)
        # Пока прежняя операция не завершилась, новая не запускается.
        with self.assertRaises(SyncError):
            sync_flow.run_with_progress(self.window, "t", lambda token: "y")
        gate.set()
        self._wait(lambda: not sync_flow._abandoned)
        self.assertEqual(sync_flow.run_with_progress(self.window, "t", lambda token: "ok"), "ok")

    def test_cancel_after_commit_waits_for_the_result(self):
        gate = threading.Event()

        def operation(token):
            token.commit()                               # необратимый шаг начат
            gate.wait(10)
            return "готово"

        self._reject_modal_after(100)
        QTimer.singleShot(400, gate.set)
        self.assertEqual(sync_flow.run_with_progress(self.window, "t", operation), "готово")
        self.assertFalse(sync_flow._abandoned)

    def test_error_from_worker_is_raised_in_caller(self):
        def boom(token):
            raise SyncError("плохо")
        with self.assertRaises(SyncError):
            sync_flow.run_with_progress(self.window, "t", boom)


class VpnHintTest(SyncUiTestCase):
    """Через VPN Яндекс.Диск обычно недоступен: пользователю нужно подсказать причину."""

    def _listening_port(self):
        listener = socket.socket()
        listener.bind(("127.0.0.1", 0))
        listener.listen(5)
        self.addCleanup(listener.close)
        return listener.getsockname()[1]

    def _closed_port(self):
        probe = socket.socket()
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
        probe.close()
        return port

    def test_internet_reachable_when_any_probe_host_answers(self):
        hosts = (("127.0.0.1", self._closed_port()), ("127.0.0.1", self._listening_port()))
        with mock.patch.object(sync_flow, "_PROBE_HOSTS", hosts):
            self.assertTrue(REAL_INTERNET_REACHABLE(timeout=1))

    def test_internet_not_reachable_when_no_probe_host_answers(self):
        hosts = (("127.0.0.1", self._closed_port()), ("127.0.0.1", self._closed_port()))
        with mock.patch.object(sync_flow, "_PROBE_HOSTS", hosts):
            self.assertFalse(REAL_INTERNET_REACHABLE(timeout=1))

    def test_probe_hosts_are_not_yandex(self):
        # Проверка связи должна отвечать на вопрос «интернет есть, а Яндекс нет?».
        self.assertTrue(all(not host.startswith("77.88.") for host, _port in sync_flow._PROBE_HOSTS))

    def test_network_error_text_mentions_vpn(self):
        text = sync_flow.describe_error(YaDiskError("x", None))
        self.assertIn("VPN", text)

    def test_server_failure_text_says_work_can_go_on(self):
        for error in (YaDiskError("x", 500), YaDiskError("x", 503), YaDiskError("x", 429)):
            text = sync_flow.describe_error(error)
            self.assertIn("работать можно", text)
            self.assertIn("выложатся", text)
        self.assertNotIn("работать можно", sync_flow.describe_error(YaDiskError("x", 507)))

    def test_other_errors_do_not_blame_vpn(self):
        for error in (YaDiskError("x", 401), YaDiskError("x", 507), YaDiskError("x", 500),
                      SyncError("плохо")):
            self.assertNotIn("VPN", sync_flow.describe_error(error))

    def _retry_dialog_text(self, has_changes, internet_ok):
        captured = []

        def fake_message_box(parent, icon, title, text):
            captured.append(text)
            return FakeBox()

        FakeBox.choice = "Повторить"
        with mock.patch.object(sync_flow, "message_box", side_effect=fake_message_box):
            retry = REAL_ASK_RETRY_OFFLINE(self.window, has_changes, internet_ok)
        self.assertTrue(retry)
        return captured[0]

    def test_retry_dialog_says_internet_is_fine_but_disk_unreachable(self):
        text = self._retry_dialog_text(True, True)
        self.assertIn("Интернет на компьютере есть, а Яндекс.Диск недоступен", text)
        self.assertIn("отключите его", text)
        self.assertIn("VPN", text)
        self.assertIn("не выложены", text)

    def test_retry_dialog_says_no_internet(self):
        text = self._retry_dialog_text(True, False)
        self.assertIn("Нет связи с интернетом", text)
        self.assertIn("VPN", text)

    def test_retry_dialog_when_connectivity_unknown(self):
        text = self._retry_dialog_text(False, None)
        self.assertIn("Не удалось связаться с Яндекс.Диском", text)
        self.assertNotIn("не выложены", text)
        self.assertIn("VPN", text)

    def test_connectivity_result_reaches_retry_dialog(self):
        self.server.offline = True
        for internet_ok in (True, False):
            with self.subTest(internet_ok=internet_ok):
                self.internet.return_value = internet_ok
                self.ask_retry.reset_mock()
                self.assertFalse(self.window.sync_now())
                self.assertIs(self.ask_retry.call_args[0][2], internet_ok)

    def test_connectivity_is_probed_only_after_a_network_failure(self):
        self.window.sync_now()
        self.internet.assert_not_called()
        self.server.offline = True
        self.window.sync_now()
        self.internet.assert_called_once()

    def test_unknown_connectivity_when_probe_is_cancelled(self):
        self.server.offline = True
        # synchronize падает сетевой ошибкой; проверку связи пользователь отменяет.
        with mock.patch.object(sync_flow, "synchronize", side_effect=YaDiskError("x", None)), \
                mock.patch.object(sync_flow, "run_with_progress", side_effect=SyncCancelled()):
            self.window.sync_now()
        self.assertIsNone(self.ask_retry.call_args[0][2])

    def test_slow_transfer_hint_appears_in_progress_dialog(self):
        with mock.patch.object(sync_flow, "_SLOW_HINT_MS", 50):
            dialog = sync_flow._BusyDialog(self.window, "Выгрузка…", CancelToken())
        dialog.show()
        self.assertFalse(dialog._slow_hint.isVisibleTo(dialog))
        self._wait(lambda: dialog._slow_hint.isVisibleTo(dialog))
        self.assertIn("VPN", dialog._slow_hint.text())
        dialog.finish()

    def test_no_slow_hint_when_transfer_is_quick(self):
        with mock.patch.object(sync_flow, "_SLOW_HINT_MS", 50):
            dialog = sync_flow._BusyDialog(self.window, "Выгрузка…", CancelToken())
        dialog.show()
        dialog.finish()
        end = time.monotonic() + 0.3
        while time.monotonic() < end:
            self.app.processEvents()
            time.sleep(0.01)
        self.assertFalse(dialog._slow_hint.isVisibleTo(dialog))


class StartupResultDeferralTest(SyncUiTestCase):
    """Результат фоновой проверки не должен всплывать поверх чужих окон."""

    def setUp(self):
        super().setUp()
        self.sync_now = mock.Mock(return_value=True)
        self.window.sync_now = self.sync_now
        self.check = SyncCheck(SyncStatus.LOCAL_AHEAD, None)

    def test_handled_immediately_when_nothing_else_is_open(self):
        self.window._handle_startup_result(None, self.check)
        self.sync_now.assert_called_once()

    def test_deferred_while_another_sync_is_running(self):
        self.window._sync_busy = True
        self.window._handle_startup_result(None, self.check)
        self.sync_now.assert_not_called()
        self.window._sync_busy = False
        self._wait(lambda: self.sync_now.called)

    def test_deferred_while_a_modal_dialog_is_open(self):
        with mock.patch("src.ui.main_window.QApplication.activeModalWidget",
                        side_effect=[object(), object(), None, None, None]):
            self.window._handle_startup_result(None, self.check)
            self.sync_now.assert_not_called()
            self._wait(lambda: self.sync_now.called)

    def test_error_window_is_also_deferred_while_busy(self):
        self.window._sync_busy = True
        self.window._handle_startup_result(YaDiskError("x", 401), None)
        self.assertEqual(self.messages, [])
        self.window._sync_busy = False
        self._wait(lambda: bool(self.messages))

    def test_ignored_once_window_is_closing(self):
        self.window._closing = True
        self.window._handle_startup_result(None, self.check)
        self.sync_now.assert_not_called()

    def test_stale_signal_without_worker_is_ignored(self):
        self.window._startup_check = None
        self.window._on_startup_check()
        self.sync_now.assert_not_called()


class UnknownChangeStateTest(SyncUiTestCase):
    """Если проверить невыложенные изменения не удалось — это не «изменений нет»."""

    def _break_local_state(self):
        with open(self.db_path + "-journal", "wb") as f:
            f.write(b"x")

    def test_indicator_says_check_failed(self):
        self._break_local_state()
        self.window.refresh_sync_indicator()
        self.assertIn("не удалось проверить", self._indicator())
        self.assertNotIn(CLEAN, self._indicator())

    def test_close_asks_when_state_is_unknown(self):
        self._break_local_state()
        self.ask_yes_no.return_value = False
        self.assertFalse(self._close(None))
        self.ask_yes_no.assert_called_once()
        self.ask_yes_no.return_value = True
        self.assertTrue(self._close(None))

    def test_offline_dialog_does_not_claim_unsynced_changes_when_unknown(self):
        self._break_local_state()
        self.server.offline = True
        self.window.sync_now()
        self.assertIs(self.ask_retry.call_args[0][1], False)


class SettingsIndicatorTest(SyncUiTestCase):
    def _log_in(self):
        """Вход по коду через локальный «Яндекс»: код сразу подтверждён."""
        self._patch_value("src.yandex_oauth.OAUTH_URL", self.server.oauth_url)
        self._patch_value("src.yandex_oauth.INFO_URL", self.server.info_url)
        self._patch("src.yandex_oauth._client_secret", return_value="test-secret")
        self.server.approved = True
        self.settings.on_login()

    def test_login_shows_sidebar_panel(self):
        save_config(SyncConfig(device_name="ПК"))
        self.settings._load()
        self.window.refresh_sync_indicator()
        self.assertFalse(self.window.sync_label.isVisibleTo(self.window))
        self._log_in()
        config = load_config()
        self.assertEqual((config.token, config.refresh_token, config.uid, config.login),
                         (self.server.token, "refresh-1", self.server.uid, self.server.login))
        self.assertGreater(config.token_expires_at, time.time())
        self.assertIn("tester", self.settings.account_label.text())
        self.assertTrue(self.window.sync_label.isVisibleTo(self.window))
        self.assertTrue(self.window.sync_btn.isVisibleTo(self.window))
        self.assertEqual(self.messages, [])

    def test_login_failure_is_shown_and_nothing_saved(self):
        save_config(SyncConfig(device_name="ПК"))
        self._patch_value("src.yandex_oauth.OAUTH_URL", self.server.url + "/missing")
        self.settings.on_login()
        self.assertTrue(any("Вход в Яндекс" == title for title, _ in self.messages))
        self.assertFalse(load_config().is_configured)

    def test_logout_hides_panel_and_forgets_token(self):
        self.window.refresh_sync_indicator()
        self.assertTrue(self.window.sync_label.isVisibleTo(self.window))
        with mock.patch("src.ui.settings_screen.ask_yes_no", return_value=True):
            self.settings.on_logout()
        self.assertFalse(self.window.sync_label.isVisibleTo(self.window))
        config = load_config()
        self.assertEqual((config.token, config.uid), ("", ""))
        self.assertEqual(config.device_name, "ПК")


class ReleaseDbTest(SyncUiTestCase):
    def test_every_screen_with_release_is_released(self):
        fake = mock.Mock()
        self.window._screens["fake"] = fake
        self.window.release_db()
        fake.release.assert_called_once()

    def test_screens_without_release_are_fine(self):
        self.window._screens["plain"] = object()
        self.window.release_db()


class RealDialogsTest(SyncUiTestCase):
    """Настоящие окна вопросов (QMessageBox): кнопки и выбор по умолчанию. Окно не
    показывается — exec_ подменён и «нажимает» нужную кнопку."""

    META = RemoteMeta(rev=2, file="f", sha256="0" * 64, schema_version=1,
                      device="Телефон", saved_at="2026-09-30 10:00:00")

    def _ask(self, function, click, *args):
        seen = {}

        def fake_exec(box):
            default, escape = box.defaultButton(), box.escapeButton()
            seen["default"] = default.text() if default else None
            seen["escape"] = escape.text() if escape else None
            seen["text"] = box.text()
            for button in box.buttons():
                if button.text() == click:
                    button.click()
                    return 0
            raise AssertionError(f"нет кнопки {click!r}: {[b.text() for b in box.buttons()]}")

        with mock.patch.object(QMessageBox, "exec_", fake_exec):
            result = function(self.window, *args)
        return result, seen

    def test_conflict_choices_and_safe_default(self):
        check = SyncCheck(SyncStatus.CONFLICT, self.META)
        for click, expected in (("Оставить эту", "local"), ("Взять с Диска", "remote"), ("Отмена", "")):
            result, seen = self._ask(REAL_ASK_CONFLICT, click, check)
            self.assertEqual(result, expected)
        self.assertEqual(seen["default"], "Отмена")      # Enter ничего не перезаписывает
        self.assertEqual(seen["escape"], "Отмена")

    def test_conflict_shows_sync_time_of_both_versions(self):
        check = SyncCheck(SyncStatus.CONFLICT, self.META, local_synced_at="2026-09-29 08:15:00")
        _, seen = self._ask(REAL_ASK_CONFLICT, "Отмена", check)
        self.assertIn("последняя синхронизация 29.09.2026 08:15", seen["text"])
        self.assertIn("«Телефон», сохранена 30.09.2026 10:00", seen["text"])

    def test_never_synced_conflict_offers_only_take_from_disk(self):
        check = SyncCheck(SyncStatus.CONFLICT, self.META, never_synced=True)
        for click, expected in (("Взять с Диска", "remote"), ("Отмена", "")):
            result, seen = self._ask(REAL_ASK_CONFLICT, click, check)
            self.assertEqual(result, expected)
        with self.assertRaises(AssertionError):            # кнопки «Оставить эту» нет
            self._ask(REAL_ASK_CONFLICT, "Оставить эту", check)
        self.assertEqual(seen["default"], "Отмена")
        self.assertIn("ни разу её не загружал", seen["text"])

    def test_conflict_with_unknown_local_sync_time(self):
        check = SyncCheck(SyncStatus.CONFLICT, self.META)
        _, seen = self._ask(REAL_ASK_CONFLICT, "Отмена", check)
        self.assertIn("последняя синхронизация —", seen["text"])

    def test_pull_question_default_is_no_when_disk_version_is_older(self):
        older = SyncCheck(SyncStatus.REMOTE_AHEAD, self.META, remote_is_older=True)
        result, seen = self._ask(REAL_ASK_PULL, "Нет", older)
        self.assertFalse(result)
        self.assertEqual(seen["default"], "Нет")

    def test_pull_question_default_is_yes_for_normal_newer_base(self):
        newer = SyncCheck(SyncStatus.REMOTE_AHEAD, self.META)
        result, seen = self._ask(REAL_ASK_PULL, "Да", newer)
        self.assertTrue(result)
        self.assertEqual(seen["default"], "Да")

    def test_empty_disk_question_default_is_no(self):
        result, seen = self._ask(REAL_ASK_EMPTY_DISK, "Да")
        self.assertTrue(result)
        self.assertEqual(seen["default"], "Нет")

    def test_retry_dialog_default_is_retry_and_later_declines(self):
        result, seen = self._ask(REAL_ASK_RETRY_OFFLINE, "Позже", True, None)
        self.assertFalse(result)
        self.assertEqual(seen["default"], "Повторить")
        self.assertEqual(seen["escape"], "Позже")
        result, _ = self._ask(REAL_ASK_RETRY_OFFLINE, "Повторить", True, True)
        self.assertTrue(result)

    def test_yes_no_helper_default_button(self):
        _, seen = self._ask(lambda w: sync_flow.ask_yes_no(w, "t", "x", default_yes=False), "Нет")
        self.assertEqual(seen["default"], "Нет")
        _, seen = self._ask(lambda w: sync_flow.ask_yes_no(w, "t", "x", default_yes=True), "Да")
        self.assertEqual(seen["default"], "Да")


class RealFlowFromScreenTest(SyncUiTestCase):
    """Экран вызывает настоящий on_data_changed: импорт → конфликт → загрузка с Диска."""

    def test_import_then_conflict_take_remote_does_not_break_screens(self):
        self.window.sync_now()
        self._push_from_other("222")
        self.ask_conflict.return_value = "remote"
        path = os.path.join(self.dir, "r.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Этикетка\nНазвание", "Номер отправления", "Статус", "Ячейка"])
        ws.append(["ii1\nТовар", "111-A-1", "Готово к выдаче", "1-1"])
        wb.save(path)
        dashboard = self.window._screens["dashboard"]
        with mock.patch("src.ui.dashboard_screen.QFileDialog.getOpenFileName", return_value=(path, "")), \
                mock.patch("src.ui.dashboard_screen.QMessageBox.information"):
            dashboard.on_import()
        self.ask_conflict.assert_called_once()
        self.assertIsNot(self.window._screens["dashboard"], dashboard)   # экраны пересозданы
        self.assertEqual(self._clients_on_screen(), ["111", "222"])
        self.app.processEvents()                                        # отложенное удаление не падает


class UntrustedDataRenderingTest(SyncUiTestCase):
    """База могла прийти с Диска: её строки нельзя показывать как разметку."""

    EVIL = '<img src="file://attacker.example/share/x.png"><b>жирный</b>'

    def test_stale_screen_shows_database_strings_as_text(self):
        from src.models import ImportSession, Shipment
        path = os.path.join(self.dir, "r.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Этикетка\nНазвание", "Номер отправления", "Статус", "Ячейка"])
        ws.append(["ii1\nТовар", "111-A-1", "Готово к выдаче", "1-1"])
        wb.save(path)
        from src.services import ImportService
        session = self.db.get_session()
        ImportService(session).process_import(path)
        for imp in session.execute(select(ImportSession)).scalars():
            imp.source_file_name = self.EVIL
        for shipment in session.execute(select(Shipment)).scalars():
            shipment.product_name = self.EVIL
        session.commit()
        session.close()

        screen = self.window._screens["stale"]
        screen.refresh()
        self.assertEqual(screen.import_info_label.textFormat(), Qt.PlainText)
        self.assertIn(self.EVIL, screen.import_info_label.text())
        tooltip = screen.table.item(0, 0).toolTip()
        self.assertNotIn("<img", tooltip)
        self.assertNotIn("<b>", tooltip)
        self.assertIn("&lt;img", tooltip)


class ScreenSignalsTest(SyncUiTestCase):
    """Экраны данных сообщают главному окну об изменении базы."""

    def setUp(self):
        super().setUp()
        self.changed = mock.Mock()
        self.window.on_data_changed = self.changed

    def _report(self, rows):
        path = os.path.join(self.dir, "report.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Этикетка\nНазвание", "Номер отправления", "Статус", "Ячейка"])
        for row in rows:
            ws.append(row)
        wb.save(path)
        return path

    def test_import_report_signals_change(self):
        dashboard = self.window._screens["dashboard"]
        path = self._report([["ii1\nТовар", "111-A-1", "Готово к выдаче", "1-1"]])
        with mock.patch("src.ui.dashboard_screen.QFileDialog.getOpenFileName", return_value=(path, "")), \
                mock.patch("src.ui.dashboard_screen.QMessageBox.information"):
            dashboard.on_import()
        self.changed.assert_called_once()

    def test_failed_import_does_not_signal_change(self):
        dashboard = self.window._screens["dashboard"]
        bad = os.path.join(self.dir, "bad.xlsx")
        with open(bad, "wb") as f:
            f.write(b"not an xlsx")
        with mock.patch("src.ui.dashboard_screen.QFileDialog.getOpenFileName", return_value=(bad, "")), \
                mock.patch("src.ui.dashboard_screen.QMessageBox.critical"):
            dashboard.on_import()
        self.changed.assert_not_called()

    def test_export_signals_change(self):
        dashboard = self.window._screens["dashboard"]
        path = self._report([["ii1\nТовар", "111-A-1", "Готово к выдаче", "1-1"]])
        with mock.patch("src.ui.dashboard_screen.QFileDialog.getOpenFileName", return_value=(path, "")), \
                mock.patch("src.ui.dashboard_screen.QMessageBox.information"):
            dashboard.on_import()
            self.changed.reset_mock()
            out = os.path.join(self.dir, "out.xlsx")
            with mock.patch("src.ui.dashboard_screen.QFileDialog.getSaveFileName", return_value=(out, "")):
                dashboard.on_export(DeliveryPoint.KOMSOMOLSKAYA_4)
        self.changed.assert_called_once()

    def test_add_client_signals_change_and_cancel_does_not(self):
        screen = self.window._screens["clients"]
        data = {"ozon_client_id": "555", "full_name": "Новый", "phone": "", "point": "KOMSOMOLSKAYA_4"}
        with mock.patch.object(clients_screen.ClientDialog, "exec_", return_value=QDialog.Rejected):
            screen.on_add_client()
        self.changed.assert_not_called()
        with mock.patch.object(clients_screen.ClientDialog, "exec_", return_value=QDialog.Accepted), \
                mock.patch.object(clients_screen.ClientDialog, "get_data", return_value=data):
            screen.on_add_client()
        self.changed.assert_called_once()
        self.assertIn("555", _client_ids(self.db_path))

    def test_delete_client_signals_change_only_when_confirmed(self):
        screen = self.window._screens["clients"]
        screen.refresh_table()
        screen.table.setCurrentCell(0, 0)
        with mock.patch("src.ui.clients_screen.ask_yes_no", return_value=False):
            screen.on_delete_client()
        self.changed.assert_not_called()
        with mock.patch("src.ui.clients_screen.ask_yes_no", return_value=True):
            screen.on_delete_client()
        self.changed.assert_called_once()

    def test_dashboard_sessions_released_and_reusable(self):
        dashboard = self.window._screens["dashboard"]
        self.window.release_db()
        session = dashboard.import_service.session
        self.assertEqual(session.execute(select(Client.ozon_client_id)).scalars().all(), ["111"])
        session.close()


if __name__ == "__main__":
    unittest.main()
