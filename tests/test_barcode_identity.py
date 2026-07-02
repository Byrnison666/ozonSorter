"""Идентификация посылки по паре «номер отправления + штрихкод».

Кейс заказчика (склад Казакова 68): в одну ячейку кладут разные посылки одного
клиента, и Ozon порой повторяет номер отправления для физически другого товара
(другой штрихкод — первая строка «Этикетка»). По одному номеру такие посылки
схлопывались в дубль: программа считала их одним товаром. Ключ стал составным
(posting_number, product_label). Плюс штрихкод выводится колонкой в выгрузку для
физической сверки на выдаче.
"""
import os
import tempfile
import unittest

import openpyxl
from sqlalchemy import select, func, text

from src.database import DatabaseManager
from src.models import Client, Shipment, DeliveryPoint
from src.services import ImportService
from src.export_service import ExportService


def _report(rows):
    """rows — список (label, posting, cell). Статус — «Готово к выдаче»."""
    path = tempfile.mktemp(suffix=".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Этикетка\nНазвание", "Номер отправления", "Статус", "Ячейка"])
    for label, posting, cell in rows:
        ws.append([f"{label}\nТовар", posting, "Готово к выдаче", cell])
    wb.save(path)
    return path


class BarcodeIdentityTest(unittest.TestCase):
    POINT = DeliveryPoint.KOMSOMOLSKAYA_4

    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        self.db = DatabaseManager(db_path=self.db_path)
        self.db.create_tables()
        self.session = self.db.get_session()
        self.session.add(Client(
            ozon_client_id="111", full_name="A",
            fixed_delivery_point=self.POINT,
        ))
        self.session.commit()
        self.imp = ImportService(self.session)
        self.exp = ExportService(self.session)
        self._files = []

    def tearDown(self):
        self.session.close()
        os.remove(self.db_path)
        for f in self._files:
            if os.path.exists(f):
                os.remove(f)

    def _import(self, rows):
        path = _report(rows)
        self._files.append(path)
        return self.imp.process_import(path)

    def _export(self, session_id):
        out = tempfile.mktemp(suffix=".xlsx")
        self._files.append(out)
        self.exp.generate_export(self.POINT, out, session_id)
        wb = openpyxl.load_workbook(out)
        return list(wb.active.iter_rows(values_only=True))

    def test_same_posting_different_barcode_are_two_shipments(self):
        # Один номер отправления, но два разных штрихкода — физически разные
        # товары в одной ячейке. Не дубль: должны стать двумя записями.
        posting = "111-0001-1"
        self._import([("ii100", posting, "A-1")])
        self._import([("ii200", posting, "A-1")])
        cnt = self.session.execute(
            select(func.count()).select_from(Shipment)
            .where(Shipment.posting_number == posting)
        ).scalar_one()
        self.assertEqual(cnt, 2, "разный штрихкод → две записи, не схлопывать в дубль")

    def test_same_posting_same_barcode_is_dedup(self):
        # Тот же номер и тот же штрихкод (повтор строки/отчёта) — это дубль.
        posting = "111-0001-1"
        self._import([("ii100", posting, "A-1")])
        self._import([("ii100", posting, "A-1")])
        cnt = self.session.execute(
            select(func.count()).select_from(Shipment)
            .where(Shipment.posting_number == posting)
        ).scalar_one()
        self.assertEqual(cnt, 1, "тот же штрихкод → одна запись")

    def test_barcode_column_in_export(self):
        # Штрихкод — колонка C выгрузки: A=описание, B=номер, C=штрихкод, D=ячейка.
        self._import([("ii16153919794", "111-0001-1", "На проверку-1")])
        last = self.session.execute(
            select(Shipment).where(Shipment.posting_number == "111-0001-1")
        ).scalar_one()
        rows = self._export(last.import_session_id)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row[1], "111-0001-1", "B — номер отправления")
        self.assertEqual(row[2], "ii16153919794", "C — штрихкод")
        self.assertEqual(row[3], "На проверку-1", "D — ячейка")

    def test_empty_label_falls_back_to_posting(self):
        # Пустая этикетка → штрихкод = номер отправления (не NULL), иначе UNIQUE
        # ловил бы NULL как «различный» и плодил дубли при повторном импорте.
        path = tempfile.mktemp(suffix=".xlsx")
        self._files.append(path)
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Этикетка\nНазвание", "Номер отправления", "Статус", "Ячейка"])
        ws.append(["", "111-0009-1", "Готово к выдаче", "A-1"])
        wb.save(path)
        self.imp.process_import(path)
        self.imp.process_import(path)  # повтор не должен упасть UNIQUE
        ship = self.session.execute(
            select(Shipment).where(Shipment.posting_number == "111-0009-1")
        ).scalar_one()
        self.assertEqual(ship.product_label, "111-0009-1")

    def test_same_posting_diff_barcode_both_exported(self):
        # Обе физически разные посылки в одной ячейке должны попасть в отгрузку.
        d1 = self._import([
            ("ii100", "111-0001-1", "A-1"),
            ("ii200", "111-0001-1", "A-1"),
        ])
        rows = self._export(d1.id)
        labels = sorted(r[2] for r in rows)
        self.assertEqual(labels, ["ii100", "ii200"])


class ShipmentSchemaMigrationTest(unittest.TestCase):
    """Rebuild боевой базы со старым UNIQUE(posting_number) на составной ключ."""

    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix=".db")
        os.close(fd)

    def tearDown(self):
        if os.path.exists(self.db_path):
            os.remove(self.db_path)

    def _make_legacy_db(self):
        # Минимальная «старая» таблица shipments с UNIQUE ровно по posting_number.
        db = DatabaseManager(db_path=self.db_path)
        with db.engine.begin() as conn:
            conn.execute(text("DROP TABLE IF EXISTS shipments"))
            conn.execute(text(
                "CREATE TABLE shipments ("
                " id INTEGER NOT NULL PRIMARY KEY,"
                " posting_number VARCHAR NOT NULL UNIQUE,"
                " client_id INTEGER,"
                " ozon_client_id_raw VARCHAR NOT NULL,"
                " product_label VARCHAR,"
                " product_name VARCHAR,"
                " ozon_type VARCHAR,"
                " ozon_status VARCHAR,"
                " cell VARCHAR,"
                " shipment_date_ozon DATETIME,"
                " is_damaged BOOLEAN,"
                " is_kty BOOLEAN,"
                " barcode VARCHAR,"
                " assignment_status VARCHAR NOT NULL,"
                " assigned_point VARCHAR,"
                " import_session_id INTEGER,"
                " first_seen_at DATETIME,"
                " last_seen_at DATETIME,"
                " shipped_to_point_at DATETIME,"
                " delivered_at DATETIME,"
                " notes TEXT"
                ")"
            ))
            conn.execute(text(
                "INSERT INTO shipments"
                " (id, posting_number, ozon_client_id_raw, product_label,"
                "  assignment_status)"
                " VALUES (1, '111-0001-1', '111', 'ii100', 'TO_SHIP')"
            ))
        return db

    def test_migration_switches_to_composite_key(self):
        db = self._make_legacy_db()
        db.create_tables()  # запускает _migrate → rebuild

        with db.engine.begin() as conn:
            # Старой строки не потеряли.
            got = conn.execute(text(
                "SELECT product_label FROM shipments WHERE posting_number='111-0001-1'"
            )).scalar_one()
            self.assertEqual(got, "ii100")
            # Тот же номер с другим штрихкодом теперь вставляется (составной ключ).
            conn.execute(text(
                "INSERT INTO shipments"
                " (posting_number, ozon_client_id_raw, product_label, assignment_status)"
                " VALUES ('111-0001-1', '111', 'ii200', 'TO_SHIP')"
            ))
            cnt = conn.execute(text(
                "SELECT COUNT(*) FROM shipments WHERE posting_number='111-0001-1'"
            )).scalar_one()
            self.assertEqual(cnt, 2)

    def test_migration_idempotent(self):
        db = self._make_legacy_db()
        db.create_tables()
        db.create_tables()  # второй прогон не должен падать/дублировать
        with db.engine.begin() as conn:
            cnt = conn.execute(text("SELECT COUNT(*) FROM shipments")).scalar_one()
            self.assertEqual(cnt, 1)


if __name__ == "__main__":
    unittest.main()
