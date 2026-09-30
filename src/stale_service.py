from dataclasses import dataclass
from datetime import datetime
from typing import Any, List, Optional
from sqlalchemy.orm import Session
from sqlalchemy import select
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .export_service import natural_key
from .models import (
    Client, Shipment, ImportSession, AssignmentStatus, DeliveryPoint
)

WARN_DAYS = 3      # с этого числа дней строка подсвечивается жёлтым
DANGER_DAYS = 7    # с этого — красным

POINT_LABELS = {
    DeliveryPoint.KOMSOMOLSKAYA_4: "Комсомольская 4",
    DeliveryPoint.KOLTSEVAYA_16: "Кольцевая 16",
}

# Колонки Excel-выгрузки, порядок фиксирован. Экран показывает сокращённый набор.
COLUMNS = [
    "Дней в базе", "Ячейка", "Номер отправления", "Штрихкод", "Товар",
    "Клиент", "Ozon ID", "Телефон", "Точка", "Впервые в отчёте", "Статус",
]

_COLUMN_WIDTHS = [12, 12, 24, 22, 50, 28, 14, 18, 18, 18, 18]

_WARN_FILL = PatternFill(start_color="FFF4D6", end_color="FFF4D6", fill_type="solid")
_DANGER_FILL = PatternFill(start_color="FFE5E5", end_color="FFE5E5", fill_type="solid")


@dataclass(frozen=True)
class StaleRow:
    days: int
    cell: str
    posting_number: str
    product_label: str
    product_name: str
    client_name: str
    ozon_client_id: str
    phone: str
    point: Optional[DeliveryPoint]
    first_seen_at: datetime
    is_returned: bool

    @property
    def point_label(self) -> str:
        return POINT_LABELS.get(self.point, "—")

    @property
    def first_seen_label(self) -> str:
        return self.first_seen_at.strftime("%d.%m.%Y")

    def as_cells(self) -> List[Any]:
        """Значения строки в порядке COLUMNS (days — int, остальное — str)."""
        return [
            self.days,
            self.cell,
            self.posting_number,
            self.product_label,
            self.product_name,
            self.client_name,
            self.ozon_client_id,
            self.phone,
            self.point_label,
            self.first_seen_label,
            "Возврат" if self.is_returned else "Готово к выдаче",
        ]


class StaleShipmentService:
    def __init__(self, db_session: Session):
        self.session = db_session

    def latest_import(self) -> Optional[ImportSession]:
        return self.session.execute(
            select(ImportSession).order_by(ImportSession.id.desc()).limit(1)
        ).scalar_one_or_none()

    def list_stale(self, point: Optional[DeliveryPoint] = None, min_days: int = 0) -> List[StaleRow]:
        latest = self.latest_import()
        if latest is None:
            return []

        # Дни считаем до даты последнего отчёта, а не до now(): иначе результат
        # менялся бы сам по себе с каждым днём без новой загрузки отчёта.
        report_date = latest.started_at.date()

        stmt = (
            select(Shipment, Client)
            .join(Client, Shipment.client_id == Client.id)
            .where(
                Shipment.last_seen_import_session_id == latest.id,
                Client.is_active == True,
                Shipment.assignment_status.notin_([
                    AssignmentStatus.EXCLUDED_NOT_OURS,
                    AssignmentStatus.EXCLUDED_KTY,
                ]),
            )
        )
        if point is not None:
            stmt = stmt.where(Shipment.assigned_point == point)

        rows: List[StaleRow] = []
        for shipment, client in self.session.execute(stmt).all():
            days = max(0, (report_date - shipment.first_seen_at.date()).days)
            if days < min_days:
                continue
            rows.append(StaleRow(
                days=days,
                cell=str(shipment.cell or ""),
                posting_number=shipment.posting_number or "",
                product_label=shipment.product_label or "",
                product_name=shipment.product_name or "",
                client_name=client.full_name or "",
                ozon_client_id=client.ozon_client_id or "",
                phone=client.phone or "",
                point=shipment.assigned_point,
                first_seen_at=shipment.first_seen_at,
                is_returned=shipment.assignment_status == AssignmentStatus.RETURNED,
            ))

        rows.sort(key=lambda r: (-r.days, natural_key(r.cell), r.posting_number))
        return rows

    @staticmethod
    def export_xlsx(rows: List[StaleRow], output_path: str) -> int:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Залежавшиеся"

        alignment = Alignment(wrap_text=True, vertical='top')
        bold = Font(bold=True)

        for col, title in enumerate(COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = bold
            cell.alignment = alignment

        for idx, row in enumerate(rows, 2):
            if row.days >= DANGER_DAYS:
                fill = _DANGER_FILL
            elif row.days >= WARN_DAYS:
                fill = _WARN_FILL
            else:
                fill = None
            for col, value in enumerate(row.as_cells(), 1):
                cell = ws.cell(row=idx, column=col, value=value)
                cell.alignment = alignment
                if fill is not None:
                    cell.fill = fill

        for col, width in enumerate(_COLUMN_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

        wb.save(output_path)
        return len(rows)
