"""Описание xlsx глазами openpyxl (как его увидит ПК): argv[1] — файл.
Печатает JSON: листы, закрепление, ширины колонок и все ячейки со значением
или стилем (значение, тип, жирность, заливка, выравнивание)."""
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(sys.argv[1])
ws = wb.active
widths = {}
for dim in ws.column_dimensions.values():
    if dim.customWidth and dim.width is not None:
        # openpyxl сливает соседние колонки одной ширины в один <col min max>.
        for idx in range(dim.min, dim.max + 1):
            widths[get_column_letter(idx)] = dim.width
cells = []
for row in ws.iter_rows():
    for c in row:
        if c.value is None and not c.has_style:
            continue
        fill = c.fill
        cells.append({
            "ref": c.coordinate,
            "value": c.value,
            "type": c.data_type,
            "bold": bool(c.font.b),
            "fill": [fill.fill_type, fill.fgColor.rgb, fill.bgColor.rgb] if fill.fill_type else None,
            "wrap": bool(c.alignment.wrap_text),
            "vertical": c.alignment.vertical,
        })
print(json.dumps({
    "sheets": wb.sheetnames,
    "title": ws.title,
    "freeze": ws.freeze_panes,
    "widths": widths,
    "cells": cells,
}, ensure_ascii=False))
