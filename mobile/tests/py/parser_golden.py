"""Эталон разбора отчёта: собрать xlsx с пограничными случаями и разобрать их
ExcelParser ПК. Печатает JSON {имя: {"rows": [...]} | {"error": "..."}}; файлы
остаются в папке argv[1] — их же разбирает телефон.
"""
import json
import os
import shutil
import sys
from datetime import datetime, time

import openpyxl
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from src.parser import ExcelParser

OUT = sys.argv[1]
NEW_HEADER = ["Этикетка\nНазвание", "Номер отправления", "Ячейка"]
OLD_HEADER = ["", "Этикетка\nНазвание", "Номер отправления", "Тип", "Статус", "Ячейка",
              "Отсчётная дата отправки", "Перевозка", "Контейнер\nШтрихкод"]


def book(rows, wb=None):
    wb = wb or openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    return wb


def cases():
    yield "new_format", book([
        [], [], [], [], [],
        NEW_HEADER,
        ["0147012251-0295-1\nПередний стабилизатор", "0147012251-0295-1", "На проверку-3"],
        ["40330930-0763-1\nФутболка", "40330930-0763-1", "235-1"],
        ["штрихкод\nТовар", "ii17574018168", "77-1"],                  # КТЯ
        [None, "224933356-0001-1", "12-3"],                            # нет этикетки
        ["  ", "224933356-0001-2", "12-4"],                            # этикетка из пробелов
        ["только-штрихкод", "224933356-0001-3", "12-5"],               # без названия
        ["abc\r\nС переводом каретки\nтретья строка", "224933356-0001-4", " 12-6 "],
        ["40330930-0763-1\nФутболка", "40330930-0763-1", "235-1"],     # дубль строки
        ["x\nЧисловая ячейка", "5550001-0001-1", 15],
        ["x\nДробная ячейка", "5550001-0001-2", 15.5],
        [" лейбл \nназвание ", "5550001-0001-3", "1-1"],
        ["x\nЦифры Unicode", "١٢٣٤٥٦-0001-1", "1-2"],
        ["x\nФормула", "5550001-0001-4", "=1+1"],
        [None, None, "без номера"],
        [None, "", "пустой номер"],
        ["x\nЧисловой номер", 777000123, "1-3"],
        ["x\nВедущие пробелы", "  0012345-7-1  ", "1-4"],
    ])

    yield "old_format", book([
        [], [],
        OLD_HEADER,
        ["", "ii\nНоски", "0107174712-0101-4", "Отправление", "Отправить на склад",
         "235-1", datetime(2026, 9, 30, 14, 5, 7), "", ""],
        ["", "x\nЧистящее", "0146744646-0115-2", "Отправление\nПовреждено",
         "Готово к выдаче", "На проверку-1", 45500.25, "П-1", "К-1"],
        ["", "y\nЧайник", "0146744646-0116-1", True, " готово к выдаче ",
         "10-1", "30.09.2026", None, None],
        ["", "z\nЛампа", "0146744646-0117-1", None, None, "На проверку-2", 45500, None, None],
        ["", "w\nСтол", "0146744646-0118-1", "Повреждено", "Вернуть продавцу",
         "11-1", 0.75, None, None],
        ["", "v\nСтул", "0146744646-0119-1", "", "", "", 30, None, None],
        ["", "u\nПолка", "0146744646-0120-1", "", "", "", 10 ** 9, None, None],
    ])

    wb = openpyxl.Workbook()
    decoy = wb.active
    decoy.title = "Первый"
    for r in [NEW_HEADER, ["a\nНе этот лист", "1111111-1-1", "1-1"]]:
        decoy.append(r)
    second = wb.create_sheet("Второй")
    for r in [NEW_HEADER, ["b\nАктивный лист", "2222222-2-2", "2-2"]]:
        second.append(r)
    wb.active = 1
    yield "active_second_sheet", wb

    yield "no_header", book([["Что-то", "другое"], ["1", "2"]])

    yield "header_variants", book(
        [[]] * 19 + [["ЭТИКЕТКА\nназвание", "Номер  отправления", "яЧейка", "Ячейка"],
                     ["q\nДве ячейки", "3333333-3-3", "левая", "правая"]]
    )

    yield "header_too_low", book(
        [[]] * 20 + [NEW_HEADER, ["q\nНиже 20-й строки", "4444444-4-4", "1-1"]]
    )

    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    book([OLD_HEADER[1:],
          ["m\nЭпоха 1904", "5555555-5-5", "Отправление", "Готово к выдаче", "1-1",
           datetime(2026, 1, 2, 3, 4, 5), None, None]], wb)
    yield "epoch_1904", wb

    yield "time_only_date", book([
        OLD_HEADER[1:],
        ["t\nТолько время", "6666666-6-6", "Отправление", "Готово к выдаче", "1-1",
         time(13, 45, 10), None, None],
    ])


# Настоящий отчёт «Остатки на складе», пересохранённый в Excel владельцем:
# пустые ячейки — <c s="5"/>, строки без номера, хвост с кодами сканера.
# Обезличен (цифры и названия товаров заменены), разметка листа — как в оригинале.
FIXTURES = {
    "excel_resaved": os.path.join(os.path.dirname(__file__), "..", "fixtures",
                                  "excel_resaved_report.xlsx"),
}


def encode(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return {"dt": str(value)}
    if isinstance(value, time):
        return {"time": str(value)}
    if isinstance(value, float):
        return {"float": repr(value)}
    return value


def sources():
    for name, wb in cases():
        path = os.path.join(OUT, f"{name}.xlsx")
        wb.save(path)
        yield name, path
    for name, src in FIXTURES.items():
        path = os.path.join(OUT, f"{name}.xlsx")
        shutil.copyfile(src, path)
        yield name, path


result = {}
for name, path in sources():
    try:
        rows = ExcelParser().parse_file(path)
        result[name] = {"rows": [{k: encode(v) for k, v in r.items()} for r in rows]}
    except ValueError as e:
        result[name] = {"error": str(e)}
print(json.dumps(result, ensure_ascii=False))
